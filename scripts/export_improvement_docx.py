#!/usr/bin/env python3
import re
import sys
from pathlib import Path

try:
    from docx import Document
except Exception as exc:  # pragma: no cover
    print(f"python-docx is required: {exc}", file=sys.stderr)
    raise SystemExit(1)


def normalize_blocks(text: str):
    cleaned = (text or '').strip()
    if not cleaned:
        return []

    blocks = re.split(r"\n\s*\n+", cleaned)
    result = []
    for block in blocks:
        block = block.strip()
        if block:
            result.append(block)
    return result if result else [cleaned]


def build_docx(input_path: str, output_path: str) -> None:
    text = Path(input_path).read_text(encoding='utf-8')
    document = Document()
    normal = document.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = None

    blocks = normalize_blocks(text)
    if not blocks:
        document.add_paragraph('Dokumen hasil perbaikan kosong.')
        document.save(output_path)
        return

    for block in blocks:
        paragraphs = [p.strip() for p in re.split(r"\n+", block)]
        for paragraph in paragraphs:
            if paragraph:
                document.add_paragraph(paragraph)
        document.add_paragraph('')

    document.save(output_path)


def main() -> int:
    if len(sys.argv) < 3:
        print('Usage: export_improvement_docx.py <input_txt> <output_docx>', file=sys.stderr)
        return 1

    input_path = sys.argv[1]
    output_path = sys.argv[2]
    build_docx(input_path, output_path)
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
