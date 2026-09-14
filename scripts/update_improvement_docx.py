#!/usr/bin/env python3
import sys
from pathlib import Path

try:
    from docx import Document
except Exception as exc:  # pragma: no cover
    print(f"python-docx is required: {exc}", file=sys.stderr)
    raise SystemExit(1)


def split_blocks(text: str):
    cleaned = (text or '').strip()
    if not cleaned:
        return []

    blocks = []
    for block in cleaned.split('\n\n'):
        value = block.strip()
        if value:
            blocks.append(value)

    return blocks or [cleaned]


def extract_replacements(content: str):
    cleaned = (content or '').strip()
    if not cleaned:
        return []

    direct_pairs = []
    lines = [line.strip() for line in cleaned.splitlines() if line.strip()]
    for line in lines:
        if '::' in line:
            old, new = [part.strip() for part in line.split('::', 1)]
            if old and new:
                direct_pairs.append((old, new))

    if direct_pairs:
        return direct_pairs

    blocks = split_blocks(cleaned)
    if len(blocks) >= 2:
        return [(blocks[i], blocks[i + 1]) for i in range(0, len(blocks) - 1, 2)]

    if len(blocks) == 1:
        return [(blocks[0], blocks[0])]

    return []


def iter_block_paragraphs(doc):
    for paragraph in doc.paragraphs:
        yield paragraph

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph

    for section in doc.sections:
        for header in (section.header, section.footer):
            if header is None:
                continue
            for paragraph in header.paragraphs:
                yield paragraph
            for table in header.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            yield paragraph


def replace_in_paragraphs(paragraphs, old, new):
    for paragraph in paragraphs:
        if old in paragraph.text:
            paragraph.text = paragraph.text.replace(old, new)


def update_docx(docx_path: str, text_path: str) -> None:
    doc = Document(docx_path)
    content = Path(text_path).read_text(encoding='utf-8')
    replacements = extract_replacements(content)

    if not replacements:
        doc.save(docx_path)
        return

    for old, new in replacements:
        if not old:
            continue

        replace_in_paragraphs(doc.paragraphs, old, new)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    replace_in_paragraphs(cell.paragraphs, old, new)

        for section in doc.sections:
            for header in (section.header, section.footer):
                if header is None:
                    continue
                replace_in_paragraphs(header.paragraphs, old, new)
                for table in header.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            replace_in_paragraphs(cell.paragraphs, old, new)

    doc.save(docx_path)


def main():
    if len(sys.argv) != 3:
        print('Usage: update_improvement_docx.py <docx_path> <text_path>', file=sys.stderr)
        return 1

    update_docx(sys.argv[1], sys.argv[2])
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
